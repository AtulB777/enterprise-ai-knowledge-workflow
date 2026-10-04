"""Text extraction for the ingestion pipeline.

Each extractor does real work against real bytes — no placeholder/fake
content per spec §55. Runs inside the background worker (see
app/workers/tasks.py), never in the request path.

CPU-bound work (PDF parsing, OCR) uses `asyncio.to_thread` so it doesn't
block the worker's event loop while a job runs.
"""

import asyncio
import csv
import io

import pytesseract
from docx import Document as DocxDocument
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader


class ExtractionError(Exception):
    def __init__(self, message: str) -> None:
        self.message = message
        super().__init__(message)


def _extract_txt_or_md(content: bytes) -> str:
    try:
        return content.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise ExtractionError("File is not valid UTF-8 text.") from exc


def _extract_csv(content: bytes) -> str:
    text = _extract_txt_or_md(content)
    # Round-trip through csv.reader mainly to fail loudly on structurally
    # broken CSV rather than silently treating garbage bytes as "text".
    try:
        rows = list(csv.reader(io.StringIO(text)))
    except csv.Error as exc:
        raise ExtractionError(f"Could not parse CSV: {exc}") from exc
    return "\n".join(", ".join(cell for cell in row) for row in rows)


def _extract_pdf_sync(content: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(content))
    except Exception as exc:  # pypdf raises several distinct error types
        raise ExtractionError(f"Could not open PDF: {exc}") from exc

    pages_text = []
    for page in reader.pages:
        pages_text.append(page.extract_text() or "")
    text = "\n\n".join(pages_text).strip()
    if not text:
        raise ExtractionError(
            "No extractable text found in PDF. It may be a scanned/image-only "
            "PDF — OCR for scanned PDFs is not yet implemented (see PLAN.md)."
        )
    return text


def _extract_docx_sync(content: bytes) -> str:
    try:
        doc = DocxDocument(io.BytesIO(content))
    except Exception as exc:
        raise ExtractionError(f"Could not open DOCX: {exc}") from exc
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    tables_text = []
    for table in doc.tables:
        for row in table.rows:
            tables_text.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(paragraphs + tables_text)


def _extract_xlsx_sync(content: bytes) -> str:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ExtractionError(f"Could not open XLSX: {exc}") from exc
    sheets_text = []
    for sheet in workbook.worksheets:
        sheets_text.append(f"# Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell) for cell in row if cell is not None]
            if cells:
                sheets_text.append(", ".join(cells))
    return "\n".join(sheets_text)


def _extract_pptx_sync(content: bytes) -> str:
    try:
        presentation = Presentation(io.BytesIO(content))
    except Exception as exc:
        raise ExtractionError(f"Could not open PPTX: {exc}") from exc
    slides_text = []
    for i, slide in enumerate(presentation.slides, start=1):
        slide_lines = [f"# Slide {i}"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs)
                    if text.strip():
                        slide_lines.append(text)
        slides_text.append("\n".join(slide_lines))
    return "\n\n".join(slides_text)


def _extract_image_ocr_sync(content: bytes) -> str:
    try:
        image = Image.open(io.BytesIO(content))
        text: str = pytesseract.image_to_string(image)
    except Exception as exc:
        raise ExtractionError(f"OCR failed: {exc}") from exc
    text = text.strip()
    if not text:
        raise ExtractionError("OCR completed but found no readable text in the image.")
    return text


_SYNC_EXTRACTORS = {
    ".pdf": _extract_pdf_sync,
    ".docx": _extract_docx_sync,
    ".xlsx": _extract_xlsx_sync,
    ".pptx": _extract_pptx_sync,
    ".png": _extract_image_ocr_sync,
    ".jpg": _extract_image_ocr_sync,
    ".jpeg": _extract_image_ocr_sync,
    ".gif": _extract_image_ocr_sync,
    ".webp": _extract_image_ocr_sync,
}


async def extract_text(*, extension: str, content: bytes) -> str:
    """Dispatches to the right extractor for `extension`. Raises
    ExtractionError with a message safe to store as the document's
    processing_error on failure.
    """
    if extension in (".txt", ".md"):
        return _extract_txt_or_md(content)
    if extension == ".csv":
        return _extract_csv(content)

    sync_extractor = _SYNC_EXTRACTORS.get(extension)
    if sync_extractor is None:
        raise ExtractionError(f"No text extractor implemented for '{extension}' yet.")
    return await asyncio.to_thread(sync_extractor, content)
