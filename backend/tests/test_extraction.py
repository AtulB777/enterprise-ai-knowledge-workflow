import io

import pytest
from docx import Document as DocxDocument
from fpdf import FPDF
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

from app.services.extraction import ExtractionError, extract_text


async def test_extracts_plain_text() -> None:
    result = await extract_text(extension=".txt", content=b"Hello, world!")
    assert result == "Hello, world!"


async def test_rejects_non_utf8_text() -> None:
    with pytest.raises(ExtractionError):
        await extract_text(extension=".txt", content=b"\xff\xfe\x00\x01invalid")


async def test_extracts_markdown_as_text() -> None:
    result = await extract_text(extension=".md", content=b"# Heading\n\nBody text.")
    assert "Heading" in result
    assert "Body text." in result


async def test_extracts_csv_content() -> None:
    csv_bytes = b"name,age\nAda,36\nGrace,85\n"
    result = await extract_text(extension=".csv", content=csv_bytes)
    assert "Ada" in result
    assert "36" in result
    assert "Grace" in result


async def test_extracts_real_pdf_text() -> None:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=14)
    pdf.cell(text="The quarterly revenue grew by twelve percent.")
    pdf_bytes = bytes(pdf.output())

    result = await extract_text(extension=".pdf", content=pdf_bytes)

    assert "quarterly revenue" in result
    assert "twelve percent" in result


async def test_corrupted_pdf_raises_extraction_error() -> None:
    with pytest.raises(ExtractionError):
        await extract_text(extension=".pdf", content=b"not actually a pdf file")


async def test_extracts_real_docx_text() -> None:
    doc = DocxDocument()
    doc.add_paragraph("Board meeting minutes for Q3.")
    doc.add_paragraph("Action item: finalize the budget.")
    buffer = io.BytesIO()
    doc.save(buffer)

    result = await extract_text(extension=".docx", content=buffer.getvalue())

    assert "Board meeting minutes" in result
    assert "finalize the budget" in result


async def test_extracts_real_xlsx_text() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Revenue"
    sheet.append(["Region", "Q1", "Q2"])
    sheet.append(["North America", 1000, 1200])
    buffer = io.BytesIO()
    workbook.save(buffer)

    result = await extract_text(extension=".xlsx", content=buffer.getvalue())

    assert "Revenue" in result
    assert "North America" in result
    assert "1200" in result


async def test_extracts_real_pptx_text() -> None:
    presentation = Presentation()
    slide_layout = presentation.slide_layouts[1]
    slide = presentation.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Company Strategy"
    buffer = io.BytesIO()
    presentation.save(buffer)

    result = await extract_text(extension=".pptx", content=buffer.getvalue())

    assert "Company Strategy" in result


async def test_ocr_extracts_text_from_image() -> None:
    image = Image.new("RGB", (600, 150), color="white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", size=48)
    draw.text((20, 40), "INVOICE PAID", fill="black", font=font)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")

    result = await extract_text(extension=".png", content=buffer.getvalue())

    assert "INVOICE" in result.upper()


async def test_ocr_on_blank_image_raises_extraction_error() -> None:
    image = Image.new("RGB", (200, 100), color="white")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")

    with pytest.raises(ExtractionError):
        await extract_text(extension=".png", content=buffer.getvalue())
