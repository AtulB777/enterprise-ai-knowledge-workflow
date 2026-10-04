import io

import pytest
from docx import Document as DocxDocument
from fpdf import FPDF

from app.core.file_validation import (
    FileValidationError,
    sanitize_display_filename,
    validate_upload,
)


def _real_pdf_bytes() -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text="Sample content")
    return bytes(pdf.output())


def _real_docx_bytes() -> bytes:
    doc = DocxDocument()
    doc.add_paragraph("Sample content")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_accepts_valid_pdf() -> None:
    extension, mime = validate_upload(
        filename="report.pdf", content=_real_pdf_bytes(), max_size_bytes=10_000_000
    )
    assert extension == ".pdf"
    assert mime == "application/pdf"


def test_accepts_valid_txt() -> None:
    extension, mime = validate_upload(
        filename="notes.txt", content=b"hello", max_size_bytes=10_000_000
    )
    assert extension == ".txt"
    assert mime == "text/plain"


def test_rejects_content_that_does_not_match_extension() -> None:
    """The core spoofing defense: a DOCX renamed to .pdf must be rejected
    even though the extension alone looks fine.
    """
    with pytest.raises(FileValidationError, match="does not match"):
        validate_upload(filename="fake.pdf", content=_real_docx_bytes(), max_size_bytes=10_000_000)


def test_rejects_disallowed_extension() -> None:
    with pytest.raises(FileValidationError, match="not supported"):
        validate_upload(filename="script.exe", content=b"MZ\x90\x00", max_size_bytes=10_000_000)


def test_rejects_oversized_file() -> None:
    with pytest.raises(FileValidationError, match="exceeds the maximum"):
        validate_upload(filename="big.txt", content=b"x" * 1000, max_size_bytes=100)


def test_rejects_empty_file() -> None:
    with pytest.raises(FileValidationError, match="empty"):
        validate_upload(filename="empty.txt", content=b"", max_size_bytes=10_000_000)


def test_sanitize_strips_path_components() -> None:
    assert sanitize_display_filename("../../etc/passwd") == "passwd"
    assert sanitize_display_filename("C:\\Windows\\evil.txt") == "evil.txt"


def test_sanitize_strips_control_characters() -> None:
    result = sanitize_display_filename("report\x00.txt")
    assert "\x00" not in result
